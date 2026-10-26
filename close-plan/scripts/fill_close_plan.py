#!/usr/bin/env python3
"""
fill_close_plan.py  -  Fill the four-tab Close Plan template from a values file.

Template-agnostic: every coordinate, bucket, marker glyph, colour and sheet name
comes from config.json, so another org can swap in their own close plan template
and declare its layout once without touching this code.

Usage:
    python fill_close_plan.py --template <template.xlsx> --values values.json \
        --out OUTPUT.xlsx [--config config.json] [--crm crm.json]

Tabs:
  1 Summary          header facts, overall score, bucket bars, top risks, next
                     actions, key people and the four deal story questions
  2 Health Score     26 items in five buckets, live per-bucket and overall formulas
  3 Action Plan      one generated row per item whose marker is not "Yes"
  4 Stakeholder Map  seniority x function heat map + persona detail table

What it does:
  - Writes ONLY what you supply. Banding, headers, widths, merged ranges, dropdowns
    and the score formulas are left exactly as the template has them.
  - Clears every mapped cell first, so an unsupplied field never reads as real
    analysis of this customer's deal.
  - Generates the Action Plan from the Health Score markers, so the two tabs cannot
    drift apart: every non-Yes item gets a row, and nothing else does.
  - Prints per-bucket and overall percentages, because openpyxl does not evaluate
    formulas and the caller needs the numbers to report them.
  - Colours every marker by CONDITIONAL FORMATTING rule rather than by a fill
    painted once at build time, so the AE's dropdown override recolours itself
    instead of leaving red text reading "Yes".
  - Writes the coverage note and warns when it is missing, because the overall
    percentage is coverage of what calls can evidence, not deal health.

values.json format (every field optional; omit what you do not want to touch):
{
  "header":   {"customer": {"value": "Acme", "sourced": true}, ...},
  "synopsis": {"business_units": {"value": "...", "sourced": true}, ...},
  "health":   {"champion": "yes", "security_review": "no", ...},
  "actions":  {"security_review": {"action": "...", "owner": "...",
                                   "status": "Not started", "comments": "..."}},
  "coverage_note": "what the score can and cannot see  (REQUIRED)",
  "summary":  {"risks": ["one line each, most serious first", ...],
               "next_actions": ["economic_buyer", "security_review", ...]},
  "stakeholders": {
    "csuite": "text",
    "grid": [{"function": "IT", "seniority": "Director",
              "name": "...", "engagement": "engaged"}],
    "personas": [{"role": "Champion", "name": "...",
                  "engagement": "engaged", "notes": "... (REQUIRED, '?' rows too)"}],
    "unplaced": "attendees the grid cannot hold, with the reason",
    "gap_line": "..."
  }
}

Requires openpyxl (pip install openpyxl --break-system-packages if missing).
"""
import argparse, datetime, json, os, re, sys

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required. Install with: pip install openpyxl --break-system-packages")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_template import summary_layout  # noqa: E402  one definition of the Summary rows


DEFAULT_COL_W = 8.43
LINE_PT = 14.5


# --------------------------------------------------------------------------- io

def load_config(explicit, template):
    here = os.path.dirname(os.path.abspath(__file__))
    for cand in [explicit,
                 os.path.join(os.path.dirname(os.path.abspath(template)), "config.json"),
                 os.path.join(here, "config.json"),
                 os.path.join(os.path.dirname(here), "config.json")]:
        if cand and os.path.exists(cand):
            with open(cand) as f:
                return json.load(f)
    sys.exit("No config.json found. Pass --config explicitly.")


def write_value(ws, coord, value, sourced, colors):
    """Write a value, black when sourced, grey italic when it is a placeholder."""
    if value is None:
        return
    cell = ws[coord]
    cell.value = value
    base = cell.font
    cell.font = Font(name=base.name, size=base.size, bold=base.bold,
                     italic=not sourced,
                     color=colors["sourced"] if sourced else colors["placeholder"])


# ------------------------------------------------------------------------- crm

CRM_LABELS = {"salesforce": "Salesforce", "hubspot": "HubSpot"}


def build_opp_link(cfg, crm_block, source=None):
    """Build the opportunity reference from whichever CRM the org actually uses.

    `source` may be 'salesforce', 'hubspot', or 'auto'. On 'auto' the first CRM with an
    opp id wins. A clickable URL needs both an opp id and a configured `instance`
    hostname; `instance` ships null, so the normal outcome is the plain text fallback
    "<CRM> - Opportunity <id>" rather than a link.

    Returns None only when there is no opp id at all. The rule the workbook enforces is
    that B2 is never blank while an opp id exists, and never a guessed hostname: an
    unlinked reference still tells the AE which record this plan was built against, and
    that is the part that matters when two records sit on one deal.
    """
    crm_cfg = cfg.get("crm") or {}
    source = source or crm_cfg.get("source") or "auto"
    order = [source] if source != "auto" else ["salesforce", "hubspot"]
    for name in order:
        spec = crm_cfg.get(name) or {}
        fields = spec.get("fields") or {}
        opp_id = (crm_block or {}).get(fields.get("opp_id", ""))
        if not opp_id:
            continue
        instance = spec.get("instance")
        if instance and spec.get("opp_url_template"):
            return spec["opp_url_template"].format(instance=instance, id=opp_id)
        label = spec.get("display_name") or CRM_LABELS.get(name) or name.title()
        return "%s - Opportunity %s" % (label, opp_id)
    return None


def resolve_crm(cfg, crm_block, source=None):
    """Resolve header values from the Aircover deal record's `crm` block.

    Makes `config.crm.fields` enforced rather than advisory: the engine reads the
    block itself instead of trusting a hand-built values.json. Returns a dict of
    {amount, close_date, stage, opp_id, opp_name, source, stage_band} with None for
    anything the block does not supply. Never invents a value.
    """
    crm_cfg = cfg.get("crm") or {}
    source = source or crm_cfg.get("source") or "auto"
    order = [source] if source != "auto" else ["salesforce", "hubspot"]
    for name in order:
        spec = crm_cfg.get(name) or {}
        fields = spec.get("fields") or {}
        got = {k: (crm_block or {}).get(v) for k, v in fields.items() if not k.startswith("_")}
        if not any(got.values()):
            continue
        stage = got.get("stage")
        tax = crm_cfg.get("stage_taxonomy") or {}
        band = None
        for b in ("early", "mid", "late", "closed", "excluded"):
            if stage and stage in (tax.get(b) or []):
                band = b
                break
        got["source"] = name
        got["stage_band"] = band
        got["opp_link"] = build_opp_link(cfg, crm_block, source=name)
        return got
    return {}


def crm_matches_deal(resolved, expected_name=None, expected_amount=None):
    """True when the CRM opportunity looks like the deal being planned.

    A deal record often carries an unrelated opportunity (a closed-won original, a
    renewal dated next year). Returns (ok, reason) so the caller can stop and ask
    rather than silently mislabelling the workbook.
    """
    if not resolved:
        return False, "CRM block is empty"
    if resolved.get("stage_band") == "closed":
        return False, "CRM opportunity is %s, a closed stage" % resolved.get("stage")
    if expected_name and resolved.get("opp_name"):
        a = str(expected_name).lower().strip()
        b = str(resolved["opp_name"]).lower().strip()
        if a not in b and b not in a:
            return False, "CRM opportunity is '%s', expected '%s'" % (resolved["opp_name"], expected_name)
    if expected_amount and resolved.get("amount"):
        try:
            if abs(float(resolved["amount"]) - float(expected_amount)) > 1:
                return False, "CRM amount %s does not match %s" % (resolved["amount"], expected_amount)
        except (TypeError, ValueError):
            pass
    return True, "match"


# ----------------------------------------------------------------------- score

def bucket_items(cfg):
    """Yield (bucket_key, bucket_spec, item_key, row) in sheet order."""
    for bkey, bspec in (cfg["health"]["buckets"]).items():
        for ikey, row in sorted(bspec["rows"].items(), key=lambda kv: kv[1]):
            yield bkey, bspec, ikey, row


def compute_scores(cfg, health):
    """Per-bucket and overall percentages, mirroring the sheet's own formulas.

    Denominator is every item in the bucket, not every answered item: an unanswered
    row counts against the score rather than vanishing from it.
    """
    per = {}
    for bkey, bspec in cfg["health"]["buckets"].items():
        rows = bspec["rows"]
        weights = {"yes": 1.0, "partial": 0.5, "no": 0.0, "unknown": 0.0}
        earned = sum(weights.get(health.get(k), 0.0) for k in rows)
        per[bkey] = 100.0 * earned / len(rows) if rows else 0.0
    overall = sum(per.values()) / len(per) if per else 0.0
    return per, overall


# --------------------------------------------------------------------- autofit

def merged_range_for(sheet, coord):
    cell = sheet[coord]
    for rng in sheet.merged_cells.ranges:
        if (rng.min_row <= cell.row <= rng.max_row
                and rng.min_col <= cell.column <= rng.max_col):
            return rng
    return None


def bulletize(text):
    """Multi-line blocks read as walls of text. Give every line a bullet so the
    content scans as a list. Leave numbered steps and existing bullets alone."""
    if not text or not isinstance(text, str) or "\n" not in text:
        return text
    out = []
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        if line[0] in "-*\u2022" or (line[:2].rstrip(".").isdigit() and "." in line[:3]):
            out.append(line)
        else:
            out.append("\u2022 " + line)
    return "\n".join(out)


def autofit(sheet, coord, max_pt=None, bullets=True):
    """Merged cells never auto-fit row height, so long narrative blocks render
    clipped no matter that wrap_text is on. Compute the wrapped line count against
    the merged range's real width and set the height explicitly."""
    cell = sheet[coord]
    text = bulletize(cell.value) if bullets else cell.value
    if not text or not isinstance(text, str):
        return
    cell.value = text
    rng = merged_range_for(sheet, coord)
    if rng:
        cols = range(rng.min_col, rng.max_col + 1)
        rows = list(range(rng.min_row, rng.max_row + 1))
    else:
        cols, rows = [cell.column], [cell.row]
    width = 0.0
    for c in cols:
        cd = sheet.column_dimensions.get(get_column_letter(c))
        width += (cd.width if cd and cd.width else DEFAULT_COL_W)
    chars = max(int(width * 1.05), 20)
    lines = 0
    for para in str(text).split("\n"):
        lines += max(1, -(-len(para) // chars))
    total = max(lines * LINE_PT, 15.0)
    if max_pt:
        total = min(total, max_pt)
    per = total / len(rows)
    for r in rows:
        existing = sheet.row_dimensions[r].height or 0
        sheet.row_dimensions[r].height = max(existing, per)
    cell.alignment = Alignment(wrap_text=True, vertical="top",
                               horizontal=cell.alignment.horizontal)


def autofit_row(sheet, row, cols, min_pt):
    """Height for a plain (unmerged) table row, from its widest wrapped cell."""
    lines = 1
    for col in cols:
        cell = sheet[f"{col}{row}"]
        if not isinstance(cell.value, str):
            continue
        cd = sheet.column_dimensions.get(col)
        chars = max(int((cd.width if cd and cd.width else DEFAULT_COL_W) * 1.05), 12)
        n = 0
        for para in cell.value.split("\n"):
            n += max(1, -(-len(para) // chars))
        lines = max(lines, n)
    sheet.row_dimensions[row].height = max(min_pt, lines * LINE_PT)


# ------------------------------------------------------------------------ tabs

def fill_overview(ws, cfg, vals, colors):
    def blank(coord):
        c = ws[coord]
        c.value = None
        base = c.font
        c.font = Font(name=base.name, size=base.size, bold=base.bold,
                      italic=True, color=colors["placeholder"])

    for key, coord in cfg["header_cells"].items():
        if not key.startswith("_"):
            blank(coord)
    for key, coord in cfg["synopsis_cells"].items():
        if not key.startswith("_"):
            blank(coord)

    for key, coord in cfg["header_cells"].items():
        if key.startswith("_"):
            continue
        item = (vals.get("header") or {}).get(key)
        if not item:
            continue
        v, sourced = item.get("value"), item.get("sourced", True)
        if key == "close_date" and v:
            try:
                v = datetime.datetime.strptime(str(v)[:10], "%Y-%m-%d")
            except ValueError:
                pass
        if key == "arr" and v not in (None, ""):
            try:
                v = float(str(v).replace(",", "").replace("$", ""))
            except ValueError:
                pass
        write_value(ws, coord, v, sourced, colors)
        # A CRM amount lands as a bare number otherwise, which reads as a quantity
        # rather than money. Currency symbol is org-specific, so it comes from config.
        fmt = (cfg.get("formats") or {}).get(key)
        if fmt and v not in (None, ""):
            ws[coord].number_format = fmt

    for key, coord in cfg["synopsis_cells"].items():
        if key.startswith("_"):
            continue
        item = (vals.get("synopsis") or {}).get(key)
        if item:
            write_value(ws, coord, item.get("value"), item.get("sourced", True), colors)
    for key, coord in cfg["synopsis_cells"].items():
        if not key.startswith("_"):
            autofit(ws, coord)


def fill_health(ws, cfg, health):
    """Write one marker per item and refresh the bucket formulas.

    The formulas are rewritten rather than trusted, so a workbook whose rows were
    edited by hand still scores against the range the config declares.
    """
    hcfg = cfg["health"]
    mcol, lcol, scol = hcfg["marker_col"], hcfg["label_col"], hcfg["score_col"]
    markers = cfg["markers"]

    for _, _, _, row in bucket_items(cfg):
        ws[f"{mcol}{row}"].value = None

    for _, _, ikey, row in bucket_items(cfg):
        status = health.get(ikey)
        if status not in markers:
            status = "unknown"
        m = markers[status]
        cell = ws[f"{mcol}{row}"]
        cell.value = m["text"] or None
        cell.font = Font(name="Calibri", size=11, bold=(status != "unknown"),
                         color=m["font"])
        if m.get("fill"):
            cell.fill = PatternFill("solid", fgColor=m["fill"])
        cell.alignment = Alignment(horizontal="center", vertical="center")

    hdr_cells = []
    for bspec in hcfg["buckets"].values():
        rows = sorted(bspec["rows"].values())
        first, last = rows[0], rows[-1]
        hdr = bspec["header_row"]
        # Each item is worth an equal share of its bucket: Yes 100%, In Progress 50%,
        # No and Needs Input 0%. Keep this in step with build_template.py, which writes
        # the same formula into the blank asset. Marker text comes from config.markers,
        # never a literal, or renaming a marker there silently zeroes the scores.
        ws[f"{scol}{hdr}"] = (f'=(COUNTIF({mcol}{first}:{mcol}{last},"{markers["yes"]["text"]}")'
                              f'+0.5*COUNTIF({mcol}{first}:{mcol}{last},'
                              f'"{markers["partial"]["text"]}"))'
                              f'/COUNTA(${lcol}${first}:${lcol}${last})')
        ws[f"{scol}{hdr}"].number_format = "0%"
        hdr_cells.append(f"{scol}{hdr}")
    ws[hcfg["overall_cell"]] = "=AVERAGE(%s)" % ",".join(hdr_cells)
    ws[hcfg["overall_cell"]].number_format = "0%"

    dd = hcfg.get("dropdown") or {}
    if dd.get("enabled"):
        set_validation(ws, [f"{mcol}{r}" for _, _, _, r in bucket_items(cfg)],
                       dd.get("options") or [])


def validate_layout(cfg, hs, sms):
    """Fail loudly on a config/template mismatch instead of writing a silent mess.

    config.json and the template asset are edited together (via build_template.py).
    When they drift, the failures are quiet and ugly: a bucket range that swallows
    the next bucket's header, or a persona list that overruns into the gap line.
    """
    errs = []
    hcfg = cfg["health"]
    headers = {b["header_row"] for b in hcfg["buckets"].values()}
    seen = {}
    for bkey, bspec in hcfg["buckets"].items():
        rows = sorted(bspec["rows"].values())
        if rows != list(range(rows[0], rows[-1] + 1)):
            errs.append(f"{bkey}: rows {rows} are not contiguous")
        if rows[0] != bspec["header_row"] + 1:
            errs.append(f"{bkey}: rows start at {rows[0]}, header is {bspec['header_row']}")
        if headers & set(rows):
            errs.append(f"{bkey}: an item row collides with a bucket header")
        for r in rows:
            if r in seen:
                errs.append(f"row {r} claimed by both {seen[r]} and {bkey}")
            seen[r] = bkey
            if not hs[f"{hcfg['label_col']}{r}"].value:
                errs.append(f"{bkey}: template has no label at {hcfg['label_col']}{r}")

    note = hcfg.get("coverage_note_cell")
    if note:
        nrow = int(re.sub(r"[^0-9]", "", note))
        last_item = max(max(b["rows"].values()) for b in hcfg["buckets"].values())
        if nrow <= last_item:
            errs.append(f"coverage_note_cell {note} sits at or above the last item row "
                        f"({last_item}). Move it below the buckets with build_template.py.")

    # The seniority axis is now something orgs are expected to edit, and every row below the
    # grid is derived from its length. Catch an axis extended in config without regenerating
    # the template, which would otherwise write grid cells over the legend and unplaced line.
    scfg = cfg["stakeholder_map"]
    g = scfg["grid"]
    grid_last = g["first_row"] + len(g["seniority"]) - 1
    unplaced = scfg.get("unplaced_cell")
    if unplaced:
        urow = int(re.sub(r"[^0-9]", "", unplaced))
        if urow != grid_last + 3:
            errs.append(f"the seniority axis ends at row {grid_last}, so build_template.py puts "
                        f"the legend at {grid_last + 2} and the unplaced line at {grid_last + 3}, "
                        f"but unplaced_cell is {unplaced}. Regenerate the template and update "
                        f"personas.first_row, unplaced_cell and gap_line_row together.")

    p = cfg["stakeholder_map"]["personas"]
    rcol = p["columns"]["role"]
    gap = cfg["stakeholder_map"].get("gap_line_row")
    for i, role in enumerate(p["roles"]):
        r = p["first_row"] + i
        if gap and r >= gap:
            errs.append(f"persona '{role}' would land on row {r}, at or past the gap line "
                        f"({gap}). Extend the template with build_template.py first.")
        elif not sms[f"{rcol}{r}"].value:
            errs.append(f"persona '{role}' has no prepared row at {rcol}{r} in the template. "
                        f"Add it with build_template.py, do not just extend config.")
    if errs:
        sys.exit("Layout does not match the template:\n  " + "\n  ".join(errs))


def set_validation(ws, cells, options):
    """Attach a single list validation to `cells`, replacing any that overlap.

    The template already carries its own dropdowns. Adding a second one over the
    same cells leaves two overlapping validations in the file, which Excel reports
    as corruption and offers to repair. LibreOffice tolerates it, so this does not
    show up in a headless render.
    """
    targets = set(cells)

    def covered(dv):
        for rng in dv.sqref.ranges:
            for r, c in rng.cells:
                if f"{get_column_letter(c)}{r}" in targets:
                    return True
        return False

    ws.data_validations.dataValidation = [
        dv for dv in ws.data_validations.dataValidation if not covered(dv)]
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options),
                        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.sqref = " ".join(cells)


def shorten_label(label):
    """Turn a Health Score row label into a short gap phrase for the Action Plan.

    Drops a trailing parenthetical and an ", e.g. ..." tail, which carry examples
    rather than the gap itself. Splitting on the first "(" instead would mangle
    labels that use parentheses mid-phrase, such as "Key risk(s) to close".
    """
    s = str(label).strip()
    s = re.sub(r"\s*\([^)]*\)\s*$", "", s)
    s = re.sub(r",\s*(e\.g\.|eg\.|for example)\s.*$", "", s, flags=re.I)
    return s.strip().rstrip("?").strip()


def fill_actions(ws, health_ws, cfg, vals):
    """Generate one row per Health Score item that is not a Yes.

    The gap text is read off the Health Score sheet rather than restated here, so
    the two tabs cannot describe the same item differently.
    """
    acfg = cfg["actions"]
    cols = acfg["columns"]
    lcol = cfg["health"]["label_col"]
    first = acfg["first_row"]

    for r in range(first, first + 300):
        for col in cols.values():
            c = ws[f"{col}{r}"]
            c.value = None
            c.fill = PatternFill(fill_type=None)
        ws.row_dimensions[r].height = None

    supplied = vals.get("actions") or {}
    missing, row, rows_by_item = [], first, {}
    for bkey, _bspec, ikey, hrow in bucket_items(cfg):
        if (vals.get("health") or {}).get(ikey) == "yes":
            continue
        entry = supplied.get(ikey) or {}
        label = health_ws[f"{lcol}{hrow}"].value or ikey
        gap = entry.get("gap") or shorten_label(label)

        ws[f"{cols['bucket']}{row}"] = acfg["bucket_titles"].get(bkey, bkey)
        ws[f"{cols['gap']}{row}"] = gap
        ws[f"{cols['action']}{row}"] = entry.get("action")
        ws[f"{cols['owner']}{row}"] = entry.get("owner")
        ws[f"{cols['status']}{row}"] = entry.get(
            "status", (acfg.get("status_dropdown") or {}).get("default"))
        ws[f"{cols['comments']}{row}"] = entry.get("comments")

        rows_by_item[ikey] = row
        tint = (acfg.get("bucket_tints") or {}).get(bkey)
        b = ws[f"{cols['bucket']}{row}"]
        b.font = Font(name="Calibri", size=9, bold=True, color="FF1B2A4A")
        b.alignment = Alignment(vertical="top", wrap_text=True)
        if tint:
            b.fill = PatternFill("solid", fgColor=tint)
        for col in (cols["gap"], cols["action"], cols["owner"],
                    cols["status"], cols["comments"]):
            ws[f"{col}{row}"].alignment = Alignment(vertical="top", wrap_text=True)
            ws[f"{col}{row}"].font = Font(name="Calibri", size=10, color="FF3D4F6F")
        ws[f"{cols['status']}{row}"].alignment = Alignment(horizontal="center", vertical="top")
        autofit_row(ws, row, list(cols.values()), acfg.get("row_height", 43.5))

        if not entry.get("action"):
            missing.append(ikey)
        row += 1

    sd = acfg.get("status_dropdown") or {}
    if sd.get("enabled") and row > first:
        set_validation(ws, [f"{cols['status']}{r}" for r in range(first, row)],
                       sd.get("options") or [])
    # Filter buttons on the header row, so an AE can pull up only their own rows.
    if row > first:
        c0 = min(cols.values())
        c1 = max(cols.values())
        ws.auto_filter.ref = f"{c0}{first - 1}:{c1}{row - 1}"
    return row - first, missing, rows_by_item


def fill_stakeholders(ws, cfg, vals):
    scfg = cfg["stakeholder_map"]
    data = vals.get("stakeholders") or {}
    eng = scfg["engagement"]
    grid = scfg["grid"]
    first_col_idx = ws[grid["first_col"] + "1"].column
    seniority, functions = grid["seniority"], grid["functions"]

    def style(coord, key, text=None):
        spec = eng.get(key) or eng["none"]
        c = ws[coord]
        c.value = text
        c.fill = PatternFill("solid", fgColor=spec["fill"])
        c.font = Font(name="Calibri", size=10, bold=bool(text), color=spec["font"])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # C-Suite band
    ws[scfg["csuite_cell"]].value = None
    if data.get("csuite"):
        c = ws[scfg["csuite_cell"]]
        c.value = data["csuite"]
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if data.get("csuite_engagement"):
            spec = eng.get(data["csuite_engagement"]) or eng["none"]
            c.fill = PatternFill("solid", fgColor=spec["fill"])
        autofit(ws, scfg["csuite_cell"], bullets=False)

    # Grid: reset to neutral, then place.
    for i in range(len(seniority)):
        for j in range(len(functions)):
            style(f"{get_column_letter(first_col_idx + j)}{grid['first_row'] + i}", "none")

    for pin in (data.get("grid") or []):
        try:
            i = seniority.index(pin["seniority"])
            j = functions.index(pin["function"])
        except (KeyError, ValueError):
            print("WARNING: grid pin skipped, unknown axis:", pin)
            continue
        key = pin.get("engagement", "engaged")
        text = pin.get("name") or (scfg.get("gap_placeholder", "?") if key == "gap" else None)
        # A gap bubble is painted by RULE on the "?" text, not statically. If it were
        # filled red here, an AE typing a name over the "?" would keep a red cell that
        # now names a mapped person. Leaving it neutral means the rule owns the colour
        # and the cell falls back to neutral the moment the "?" goes away.
        style(f"{get_column_letter(first_col_idx + j)}{grid['first_row'] + i}",
              "none" if key == "gap" else key, text)

    # Persona detail
    p = scfg["personas"]
    pcols, prow0 = p["columns"], p["first_row"]
    by_role = {str(x.get("role", "")).lower(): x for x in (data.get("personas") or [])}
    noteless = []
    for i, role in enumerate(p["roles"]):
        r = prow0 + i
        rec = by_role.get(role.lower(), {})
        if not str(rec.get("notes") or "").strip():
            noteless.append(role)
        ws[f"{pcols['name']}{r}"] = rec.get("name") or scfg.get("gap_placeholder", "?")
        ws[f"{pcols['name']}{r}"].alignment = Alignment(vertical="top", wrap_text=True)
        ws[f"{pcols['notes']}{r}"] = rec.get("notes")
        ws[f"{pcols['notes']}{r}"].alignment = Alignment(vertical="top", wrap_text=True)
        key = rec.get("engagement") or "gap"
        spec = eng.get(key) or eng["gap"]
        c = ws[f"{pcols['engagement']}{r}"]
        c.value = spec["text"] or None
        c.fill = PatternFill("solid", fgColor=spec["fill"])
        c.font = Font(name="Calibri", size=10, color=spec["font"])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        autofit(ws, f"{pcols['notes']}{r}", bullets=False)

    # Attendees the grid cannot hold. The axes demand a function and a seniority, and a
    # person who joined one call with no role stated anywhere has neither. Dropping them
    # makes the committee look narrower than it is, so they are named here with the
    # reason rather than guessed onto a cell.
    up = scfg.get("unplaced_cell")
    if up:
        cell = ws[up]
        cell.value = None
        text = (data.get("unplaced") or "").strip()
        if text:
            cell.value = text
            cell.font = Font(name="Calibri", size=9, italic=True, color="FF7A8CA0")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            autofit(ws, up, bullets=False)

    # The template carries the standing open question. Append the answer to it
    # rather than overwriting, so an unanswered workbook still asks the question.
    gap_row = scfg.get("gap_line_row")
    if gap_row:
        cell = ws[f"A{gap_row}"]
        prompt = str(cell.value or "").strip()
        answer = (data.get("gap_line") or "").strip()
        cell.value = (prompt + " " + answer).strip() if prompt else (answer or None)
        if cell.value:
            autofit(ws, f"A{gap_row}", bullets=False)

    mapped = sum(1 for x in (data.get("personas") or [])
                 if x.get("name") and x.get("engagement") in ("engaged", "emerging"))
    # A grid where every attendee fell through to the unplaced line used to ship silently.
    # It nearly always means a source was skipped (an attendee title in the transcript header,
    # or a Stakeholder Coverage role line), not that the committee is genuinely unmappable.
    named_cells = sum(1 for pin in (data.get("grid") or []) if pin.get("name"))
    return mapped, len(p["roles"]), named_cells, noteless


# --------------------------------------------------------------------- summary

def fill_summary(ws, cfg, vals, rows_by_item):
    """Top risks and next actions on the Summary tab.

    Risks are written as text. Next actions are FORMULAS into the Action Plan rows,
    so an AE who rewrites an action or moves its status sees the Summary follow.
    Only the first sentence of each action shows here; the exit criterion stays on
    the Action Plan, which keeps the Summary to one page.
    """
    lay = summary_layout(cfg)
    aps = cfg["template"]["actions_sheet"]
    acols = cfg["actions"]["columns"]
    data = vals.get("summary") or {}
    health = vals.get("health") or {}
    notes = []

    risks = [str(r).strip() for r in (data.get("risks") or []) if str(r).strip()]
    if not risks:
        notes.append("summary.risks not supplied; listed the first items scored No instead")
        hs = ws.parent[cfg["template"]["health_sheet"]]
        lcol = cfg["health"]["label_col"]
        # A bare checklist label ("Procurement process started") reads as good news,
        # so the fallback says plainly that the item is missing.
        risks = ["Not in place: " + shorten_label(hs[f"{lcol}{hrow}"].value)
                 for _, _, ikey, hrow in bucket_items(cfg) if health.get(ikey) == "no"]
    elif len(risks) > len(lay["risk_rows"]):
        notes.append("summary.risks has %d entries; the Summary shows the first %d"
                     % (len(risks), len(lay["risk_rows"])))
    for row in lay["risk_rows"]:
        ws[f"B{row}"].value = None
    for row, text in zip(lay["risk_rows"], risks):
        ws[f"B{row}"].value = "\u2022 " + text
        autofit(ws, f"B{row}", bullets=False)

    order = [k for k in (data.get("next_actions") or []) if k in rows_by_item]
    unknown = [k for k in (data.get("next_actions") or []) if k not in rows_by_item]
    if unknown:
        notes.append("summary.next_actions names items with no Action Plan row (scored Yes "
                     "or misspelled): " + ", ".join(unknown))
    if not order:
        notes.append("summary.next_actions not supplied; took No items first, then the rest")
        rank = {"no": 0, "unknown": 1, "partial": 2}
        order = sorted(rows_by_item, key=lambda k: (rank.get(health.get(k), 3),
                                                    rows_by_item[k]))
    width = sum(ws.column_dimensions[c].width or DEFAULT_COL_W for c in "CDEF")
    for i, row in enumerate(lay["action_rows"]):
        for c in ("B", "C", "G", "I"):
            ws[f"{c}{row}"].value = None
        if i >= len(order):
            continue
        arow = rows_by_item[order[i]]
        src = f"'{aps}'!{acols['action']}{arow}"
        ws[f"B{row}"] = i + 1
        ws[f"C{row}"] = f'=IFERROR(LEFT({src},FIND(". ",{src})),{src})'
        ws[f"G{row}"] = f"='{aps}'!{acols['owner']}{arow}"
        ws[f"I{row}"] = f"='{aps}'!{acols['status']}{arow}"
        action = str((vals.get("actions") or {}).get(order[i], {}).get("action") or "")
        first = action.split(". ")[0]
        lines = max(1, -(-len(first) // max(int(width * 1.2), 20)))
        ws.row_dimensions[row].height = max(22, lines * LINE_PT + 6)
    return notes


# ------------------------------------------------------------- coverage note

def fill_coverage_note(cfg, hs, vals):
    """Write the coverage note, or warn loudly that it is missing.

    Transcribed calls skew toward technical working sessions and usually stop before
    close, so the overall percentage measures what the transcript corpus can
    evidence rather than the health of the deal. An AE who reads 10% as "this is
    dying" when the pricing window was simply never transcribed has been actively
    misled by the workbook. Stating the limit on the face of the sheet is the
    cheapest fix there is, so this cell is required rather than optional.
    """
    hcfg = cfg["health"]
    coord = hcfg.get("coverage_note_cell")
    if not coord:
        return False
    cell = hs[coord]
    cell.value = None
    note = (vals.get("coverage_note") or "").strip()
    if not note:
        return False
    cell.value = note
    cell.font = Font(name="Calibri", size=9, color=cfg["value_colors"]["placeholder"])
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    autofit(hs, coord, bullets=False)
    return True


# ------------------------------------------------------------ live colouring

def add_live_colors(cfg, hs, aps, sms, sus=None):
    """Colour every marker by RULE, not by a fill painted at build time.

    A static PatternFill is a snapshot. The AE picks a new value from the dropdown,
    the bucket formula recomputes and the row keeps its old colour, so the sheet
    shows red next to the word "Yes". That does not read as a stale fill, it reads
    as a broken score, and it undermines every other number on the tab.
    Conditional formatting keyed on the marker text moves with the value and
    survives the Excel and Google Sheets round trip.
    """
    from openpyxl.formatting.rule import Rule
    from openpyxl.styles.differential import DifferentialStyle

    def rules(ws, ranges, mapping):
        for text, spec in mapping.items():
            if not text:
                continue
            fill = spec.get("fill")
            dxf = DifferentialStyle(
                fill=PatternFill(start_color=fill, end_color=fill, fill_type="solid")
                if fill else None,
                font=Font(color=spec.get("font"), bold=spec.get("bold", False)))
            r = Rule(type="cellIs", operator="equal", dxf=dxf, stopIfTrue=False)
            r.formula = ['"%s"' % text]
            for rng in ranges:
                ws.conditional_formatting.add(rng, r)

    def blank_rule(ws, rng, spec):
        fill = spec.get("fill")
        dxf = DifferentialStyle(
            fill=PatternFill(start_color=fill, end_color=fill, fill_type="solid"))
        ws.conditional_formatting.add(
            rng, Rule(type="containsBlanks", dxf=dxf, stopIfTrue=False))

    # Health Score markers, one contiguous range per bucket.
    m, hcfg = cfg["markers"], cfg["health"]
    mcol = hcfg["marker_col"]
    ranges = []
    for bspec in hcfg["buckets"].values():
        rws = sorted(bspec["rows"].values())
        ranges.append(f"{mcol}{rws[0]}:{mcol}{rws[-1]}")
    rules(hs, ranges, {m[k]["text"]: dict(m[k], bold=(k != "unknown"))
                       for k in ("yes", "partial", "no", "unknown")})

    scfg = cfg["stakeholder_map"]
    eng = scfg["engagement"]

    # Persona detail: the Engagement column is a dropdown, so it is fully rule-driven.
    p = scfg["personas"]
    ecol = p["columns"]["engagement"]
    first, last = p["first_row"], p["first_row"] + len(p["roles"]) - 1
    rules(sms, [f"{ecol}{first}:{ecol}{last}"],
          {eng[k]["text"]: eng[k] for k in ("engaged", "emerging", "gap")})

    # Heat map: engagement level IS the colour here, so green and amber cannot be
    # inferred from a name and stay painted at build time. Only two states are
    # rule-driven: a "?" gap bubble, and a cell someone has emptied.
    grid = scfg["grid"]
    c0 = sms[grid["first_col"] + "1"].column
    gc0 = get_column_letter(c0)
    gc1 = get_column_letter(c0 + len(grid["functions"]) - 1)
    gr0 = grid["first_row"]
    gr1 = gr0 + len(grid["seniority"]) - 1
    grng = f"{gc0}{gr0}:{gc1}{gr1}"
    rules(sms, [grng], {scfg.get("gap_placeholder", "?"): eng["gap"]})
    blank_rule(sms, grng, eng["none"])

    # Action Plan status column, palette from config so it is not engine-owned.
    acfg = cfg["actions"]
    sd = acfg.get("status_dropdown") or {}
    palette = sd.get("colors") or {}
    if sd.get("enabled") and palette:
        scol = acfg["columns"]["status"]
        gcol = acfg["columns"]["gap"]
        rows = [r for r in range(acfg["first_row"], acfg["first_row"] + 300)
                if aps[f"{gcol}{r}"].value]
        if rows:
            rules(aps, [f"{scol}{rows[0]}:{scol}{rows[-1]}"], palette)

    # Summary: the status and engagement cells are formulas, so they need the same
    # rules as their sources or they would show plain text.
    if sus is not None:
        lay = summary_layout(cfg)
        a0, a1 = lay["action_rows"][0], lay["action_rows"][-1]
        if palette:
            rules(sus, [f"I{a0}:I{a1}"], palette)
        p0, p1 = lay["people_rows"][0], lay["people_rows"][-1]
        rules(sus, [f"H{p0}:H{p1}"],
              {eng[k]["text"]: eng[k] for k in ("engaged", "emerging", "gap")})


def wrap_health_labels(cfg, hs):
    """Wrap and fit the item labels.

    Several are long enough to clip at the column width, and a clipped label such
    as "Have we established enough value w/ key stakeholders to mitigate potential
    for" changes what the item appears to ask.
    """
    hcfg = cfg["health"]
    lcol = hcfg["label_col"]
    for _, _, _, row in bucket_items(cfg):
        c = hs[f"{lcol}{row}"]
        c.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal=c.alignment.horizontal)
        autofit_row(hs, row, [lcol], 15.0)


# ------------------------------------------------------------------ solid bars

X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"


def solid_databars(path, sheet_index):
    """Turn the Summary's data bars from gradient to solid, after openpyxl has saved.

    openpyxl only writes the Excel 2007 data bar, which fades to white and never reaches
    zero length. Excel 2010+ reads an x14 extension for a solid, zero-based bar. openpyxl
    drops that extension on load, so it is added to the saved file here, in the same
    XML Excel itself writes. Readers without x14 support fall back to the gradient bar.
    """
    import uuid
    import zipfile

    name = f"xl/worksheets/sheet{sheet_index}.xml"
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    xml = dict((i.filename, d) for i, d in items)[name].decode("utf-8")

    rule_re = re.compile(r'(<conditionalFormatting sqref="([^"]+)"><cfRule type="dataBar"[^>]*>'
                         r'<dataBar[^>]*>.*?</dataBar>)(</cfRule>)')
    exts = []

    def patch(m):
        gid = "{%s}" % str(uuid.uuid4()).upper()
        body = re.sub(r'<color rgb="00', '<color rgb="FF', m.group(1))
        exts.append(
            f'<x14:conditionalFormatting xmlns:xm="{XM}"><x14:cfRule type="dataBar" id="{gid}">'
            '<x14:dataBar minLength="0" maxLength="100" gradient="0">'
            '<x14:cfvo type="num"><xm:f>0</xm:f></x14:cfvo>'
            '<x14:cfvo type="num"><xm:f>1</xm:f></x14:cfvo>'
            '<x14:negativeFillColor rgb="FFD94F4F"/><x14:axisColor rgb="FF000000"/>'
            f'</x14:dataBar></x14:cfRule><xm:sqref>{m.group(2)}</xm:sqref></x14:conditionalFormatting>')
        return (body + f'<extLst><ext uri="{{B025F937-C7B1-47D3-B67F-A62EFF666E3E}}" '
                f'xmlns:x14="{X14}"><x14:id>{gid}</x14:id></ext></extLst>' + m.group(3))

    xml = rule_re.sub(patch, xml)
    if not exts or "<extLst><ext uri=\"{78C0D931" in xml:
        return
    xml = xml.replace("</worksheet>",
                      f'<extLst><ext uri="{{78C0D931-6437-407d-A8EE-F0AAD7539E65}}" xmlns:x14="{X14}">'
                      f'<x14:conditionalFormattings>{"".join(exts)}</x14:conditionalFormattings>'
                      '</ext></extLst></worksheet>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            z.writestr(info, xml.encode("utf-8") if info.filename == name else data)


# ------------------------------------------------------------------------ main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--values", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--config", default=None)
    ap.add_argument("--crm", default=None,
                    help="Optional JSON file holding the deal record's `crm` block. When given, the "
                         "engine resolves amount, close date, stage, and opp link from it through "
                         "config.crm.fields, filling only header keys the values file omitted.")
    a = ap.parse_args()

    cfg = load_config(a.config, a.template)
    with open(a.values) as f:
        vals = json.load(f)

    # Engine-side header resolution. Values supplied explicitly always win, so a
    # person's correction in the confirmation step is never overwritten by the CRM.
    if a.crm:
        with open(a.crm) as f:
            crm_block = json.load(f)
        resolved = resolve_crm(cfg, crm_block)
        hdr = vals.setdefault("header", {})
        for key, rk in (("arr", "amount"), ("close_date", "close_date"),
                        ("forecast", "stage"), ("opp_link", "opp_link")):
            if key not in hdr and resolved.get(rk) is not None:
                hdr[key] = {"value": resolved[rk], "sourced": True}
        if resolved.get("stage_band") == "closed":
            print("WARNING: CRM opportunity '%s' is at %s, a closed stage. Confirm this is the "
                  "opportunity being planned." % (resolved.get("opp_name"), resolved.get("stage")))

    colors = cfg["value_colors"]
    wb = load_workbook(a.template)
    t = cfg["template"]
    ws = wb[t["overview_sheet"]]
    hs = wb[t["health_sheet"]]
    aps = wb[t["actions_sheet"]]
    sms = wb[t["stakeholder_sheet"]]

    # The shipped template may carry a specific company's wording in its labels.
    # Rewrite them from config so nothing vendor-specific is baked into the asset.
    vendor = (cfg.get("vendor") or {}).get("name", "") or ""
    lo = cfg.get("label_overrides") or {}
    target = {"overview": ws, "health": hs, "actions": aps,
              "stakeholder": sms}.get(lo.get("sheet"), ws)
    for coord, text in (lo.get("cells") or {}).items():
        if coord.startswith("_"):
            continue
        # With no vendor configured, a {vendor} label would render as "using ?".
        # Leave the template's own neutral wording in place instead.
        if "{vendor}" in text and not vendor.strip():
            continue
        target[coord] = text.replace("{vendor}", vendor)

    validate_layout(cfg, hs, sms)

    fill_overview(ws, cfg, vals, colors)

    health = vals.get("health") or {}
    fill_health(hs, cfg, health)
    n_actions, missing, rows_by_item = fill_actions(aps, hs, cfg, vals)
    summary_notes = fill_summary(ws, cfg, vals, rows_by_item)
    mapped, total_personas, named_cells, noteless = fill_stakeholders(sms, cfg, vals)

    have_note = fill_coverage_note(cfg, hs, vals)

    add_live_colors(cfg, hs, aps, sms, ws)

    wrap_health_labels(cfg, hs)

    wb.save(a.out)
    solid_databars(a.out, wb.sheetnames.index(t["overview_sheet"]) + 1)

    per, overall = compute_scores(cfg, health)
    print("wrote", a.out)
    for bkey, pct in per.items():
        print("  %-24s %3.0f%%" % (cfg["actions"]["bucket_titles"].get(bkey, bkey), pct))
    print("  %-24s %3.0f%%" % ("OVERALL", overall))
    print("  action rows: %d   personas engaged or emerging: %d of %d   grid names: %d"
          % (n_actions, mapped, total_personas, named_cells))
    if noteless:
        print("WARNING: %d persona row(s) shipped with an empty Notes / Gap cell: %s. A '?' row "
              "needs a note more than a filled one does: say who is missing, why it matters, and "
              "how to reach them. An empty cell reads as an unfinished workbook."
              % (len(noteless), ", ".join(noteless)))
    if not named_cells:
        print("WARNING: the Stakeholder Map grid has no named placements. Every attendee fell "
              "through to the unplaced line, which usually means a source was skipped: check "
              "the attendee title in each transcript header and the Close Plan Stakeholder "
              "Coverage role lines. Where a function is evidenced but a level is not, use the "
              "'Seniority not stated' row instead of the unplaced line.")
    if missing:
        print("WARNING: no recommended action supplied for: " + ", ".join(missing))
    for note in summary_notes:
        print("NOTE (Summary tab): " + note)
    if not have_note:
        print("WARNING: no coverage_note supplied. The overall percentage is coverage of "
              "what the calls in scope can evidence, not deal health, and the workbook has "
              "to say so on its face. Treat this as a build failure.")


if __name__ == "__main__":
    main()
